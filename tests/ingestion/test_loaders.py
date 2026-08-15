from __future__ import annotations

import struct
import zipfile
from pathlib import Path

import pdfplumber
import pytest
from docx import Document as WordDocument

from docchrono.domain import CaseConfig, DocumentFailureCode
from docchrono.ingestion import DocxLoader, EmlLoader, Ingestor, PdfLoader


def test_docx_preserves_paragraphs_and_table_cells(tmp_path: Path) -> None:
    source = tmp_path / "document.docx"
    authored = WordDocument()
    authored.add_paragraph("First sentence. Second sentence!")
    table = authored.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    authored.save(str(source))

    result = Ingestor().ingest(source)

    assert not result.failures
    document = result.documents[0]
    assert document.parser_name == DocxLoader.name
    assert "First sentence." in document.raw_text
    assert [segment.kind for segment in document.segments].count("sentence") >= 2
    cells = [segment for segment in document.segments if segment.kind == "cell"]
    assert [document.raw_text[cell.raw_start : cell.raw_end] for cell in cells] == ["A1", "B1"]


def test_eml_preserves_headers_html_fallback_and_attachment(tmp_path: Path) -> None:
    source = tmp_path / "message.eml"
    source.write_bytes(
        b"From: Alice <alice@example.com>\r\n"
        b"To: Bob <bob@example.com>\r\n"
        b"Date: Fri, 14 Aug 2026 10:30:00 -0500\r\n"
        b"Subject: Test\r\n"
        b"MIME-Version: 1.0\r\n"
        b"Content-Type: multipart/mixed; boundary=outer\r\n\r\n"
        b"--outer\r\nContent-Type: text/html; charset=utf-8\r\n\r\n"
        b"<html><body><p>Hello <b>Bob</b>.</p><script>ignore()</script></body></html>\r\n"
        b"--outer\r\nContent-Type: text/plain; charset=utf-8\r\n"
        b"Content-Disposition: attachment; filename=notes.txt\r\n\r\n"
        b"Attached evidence.\r\n--outer--\r\n"
    )

    result = Ingestor().ingest(source)

    assert not result.failures
    document = result.documents[0]
    assert document.parser_name == EmlLoader.name
    assert document.metadata["headers"]["from"] == ("Alice <alice@example.com>",)
    assert document.metadata["headers"]["date"] == ("Fri, 14 Aug 2026 10:30:00 -0500",)
    assert "Hello" in document.raw_text
    assert "ignore()" not in document.raw_text
    assert "Attached evidence." in document.raw_text
    assert document.metadata["attachments"][0]["included_text"] is True
    fields = {segment.field for segment in document.segments if segment.kind == "field"}
    assert {"from", "to", "date", "subject", "body"} <= fields
    assert any(segment.kind == "attachment" for segment in document.segments)


class _FakePage:
    def __init__(self, text: str, *, images: tuple[object, ...] = ()) -> None:
        self._text = text
        self.images = images

    def extract_text(self) -> str:
        return self._text

    def extract_words(self) -> list[dict[str, object]]:
        if not self._text:
            return []
        return [{"text": "Evidence", "x0": 10, "top": 20, "x1": 50, "bottom": 30}]


class _FakePdf:
    def __init__(self, pages: list[_FakePage]) -> None:
        self.pages = pages

    def __enter__(self) -> _FakePdf:
        return self

    def __exit__(self, *args: object) -> None:
        return None


def test_pdf_preserves_pages_and_word_boxes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "document.pdf"
    source.write_bytes(b"%PDF authored fixture")

    def open_pdf(_stream: object) -> _FakePdf:
        return _FakePdf([_FakePage("Evidence here.")])

    monkeypatch.setattr(pdfplumber, "open", open_pdf)

    result = Ingestor().ingest(source)

    assert not result.failures
    document = result.documents[0]
    assert document.parser_name == PdfLoader.name
    assert document.segments[0].kind == "page"
    assert document.metadata["word_boxes"] == (
        {
            "raw_start": 0,
            "raw_end": 8,
            "page": 1,
            "x0": 10.0,
            "y0": 20.0,
            "x1": 50.0,
            "y1": 30.0,
            "text": "Evidence",
        },
    )


def test_image_only_pdf_returns_ocr_required(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"%PDF authored image fixture")

    def open_pdf(_stream: object) -> _FakePdf:
        return _FakePdf([_FakePage("", images=({},))])

    monkeypatch.setattr(pdfplumber, "open", open_pdf)

    result = Ingestor().ingest(source)

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.OCR_REQUIRED


def test_pdf_page_limit_is_typed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "long.pdf"
    source.write_bytes(b"%PDF authored long fixture")

    def open_pdf(_stream: object) -> _FakePdf:
        return _FakePdf([_FakePage("one"), _FakePage("two")])

    monkeypatch.setattr(pdfplumber, "open", open_pdf)

    result = Ingestor().ingest(source, CaseConfig(max_pdf_pages=1))

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.LIMIT_EXCEEDED


def test_pdf_extracted_character_limit_is_incremental(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "expanding.pdf"
    source.write_bytes(b"%PDF authored expansion fixture")

    def open_pdf(_stream: object) -> _FakePdf:
        return _FakePdf([_FakePage("1234"), _FakePage("x")])

    monkeypatch.setattr(pdfplumber, "open", open_pdf)

    result = Ingestor().ingest(source, CaseConfig(max_extracted_chars=5))

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.LIMIT_EXCEEDED
    assert "page 2" in result.failures[0].message


def test_pdf_word_box_limit_is_typed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    class ManyWordsPage(_FakePage):
        def extract_words(self) -> list[dict[str, object]]:
            return [
                {"text": "A", "x0": 0, "top": 0, "x1": 1, "bottom": 1},
                {"text": "B", "x0": 2, "top": 0, "x1": 3, "bottom": 1},
            ]

    source = tmp_path / "many-words.pdf"
    source.write_bytes(b"%PDF authored word fixture")

    def open_pdf(_stream: object) -> _FakePdf:
        return _FakePdf([ManyWordsPage("A B")])

    monkeypatch.setattr(pdfplumber, "open", open_pdf)

    result = Ingestor().ingest(source, CaseConfig(max_pdf_word_boxes=1))

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.LIMIT_EXCEEDED
    assert "word boxes" in result.failures[0].message


def test_encrypted_docx_signature_is_typed(tmp_path: Path) -> None:
    source = tmp_path / "encrypted.docx"
    source.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1encrypted")

    result = Ingestor().ingest(source)

    assert result.failures[0].code is DocumentFailureCode.ENCRYPTED


def test_docx_member_count_is_preflighted_before_parser(tmp_path: Path) -> None:
    source = tmp_path / "too-many-members.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "document")

    result = Ingestor().ingest(source, CaseConfig(max_docx_members=1))

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.LIMIT_EXCEEDED
    assert "archive members" in result.failures[0].message


def test_docx_member_and_total_uncompressed_limits_are_preflighted(tmp_path: Path) -> None:
    source = tmp_path / "large-member.docx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("[Content_Types].xml", "12345")
        archive.writestr("word/document.xml", "67890")

    member_result = Ingestor().ingest(source, CaseConfig(max_docx_member_bytes=4))
    total_result = Ingestor().ingest(
        source,
        CaseConfig(max_docx_member_bytes=10, max_docx_uncompressed_bytes=9),
    )

    assert member_result.failures[0].code is DocumentFailureCode.LIMIT_EXCEEDED
    assert "member" in member_result.failures[0].message
    assert total_result.failures[0].code is DocumentFailureCode.LIMIT_EXCEEDED
    assert "uncompressed archive bytes" in total_result.failures[0].message


def test_docx_high_compression_ratio_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "compression-bomb.docx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "A" * 10_000)

    result = Ingestor().ingest(source, CaseConfig(max_docx_compression_ratio=2))

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.LIMIT_EXCEEDED
    assert "compression ratio" in result.failures[0].message


def test_docx_encrypted_member_flag_is_typed(tmp_path: Path) -> None:
    source = tmp_path / "encrypted-member.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("word/document.xml", "content")
    content = bytearray(source.read_bytes())
    local_header = content.index(b"PK\x03\x04")
    central_header = content.index(b"PK\x01\x02")
    struct.pack_into("<H", content, local_header + 6, 1)
    struct.pack_into("<H", content, central_header + 8, 1)
    source.write_bytes(content)

    result = Ingestor().ingest(source)

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.ENCRYPTED


def test_docx_unsafe_member_name_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "unsafe-name.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("../word/document.xml", "content")

    result = Ingestor().ingest(source)

    assert result.documents == ()
    assert result.failures[0].code is DocumentFailureCode.MALFORMED
    assert "unsafe archive member" in result.failures[0].message
