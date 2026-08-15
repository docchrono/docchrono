from __future__ import annotations

import hashlib
import io
import re
import zipfile
from email import policy
from email.message import Message
from email.parser import BytesParser
from pathlib import PurePosixPath
from typing import Any

from docchrono.domain import DocumentFailureCode, DocumentSegment

from .contracts import LoadedDocument, LoaderFailure, LoadRequest

_OLE_COMPOUND_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_SEMANTIC_HEADERS = ("from", "to", "cc", "bcc", "date", "subject", "message-id")


class TextLoader:
    name = "text"
    version = "1.0.0"
    supported_suffixes = (".txt",)
    media_type = "text/plain"

    def load(self, request: LoadRequest) -> LoadedDocument:
        raw_text = _decode_source_text(request.content, request.path.name)
        segments, _, _ = _text_segments(raw_text)
        return LoadedDocument(
            raw_text=raw_text,
            media_type=self.media_type,
            segments=tuple(segments),
            metadata={"format": "text"},
        )


class MarkdownLoader:
    name = "markdown"
    version = "1.0.0"
    supported_suffixes = (".md", ".markdown")
    media_type = "text/markdown"

    def load(self, request: LoadRequest) -> LoadedDocument:
        raw_text = _decode_source_text(request.content, request.path.name)
        segments, _, _ = _text_segments(raw_text)
        heading_pattern = re.compile(r"(?m)^(?P<marker>#{1,6})[ \t]+(?P<title>[^\r\n]+)")
        for match in heading_pattern.finditer(raw_text):
            segments.append(
                DocumentSegment(
                    kind="field",
                    field="heading",
                    label=str(len(match.group("marker"))),
                    raw_start=match.start("title"),
                    raw_end=match.end("title"),
                )
            )
        return LoadedDocument(
            raw_text=raw_text,
            media_type=self.media_type,
            segments=tuple(_ordered_segments(segments)),
            metadata={"format": "markdown"},
        )


class PdfLoader:
    name = "pdfplumber"
    version = "1.1.0"
    supported_suffixes = (".pdf",)
    media_type = "application/pdf"

    def load(self, request: LoadRequest) -> LoadedDocument:
        try:
            import pdfplumber

            pdf_context = pdfplumber.open(io.BytesIO(request.content))
            with pdf_context as pdf:
                if len(pdf.pages) > request.config.max_pdf_pages:
                    raise LoaderFailure(
                        DocumentFailureCode.LIMIT_EXCEEDED,
                        f"PDF has {len(pdf.pages)} pages; limit is {request.config.max_pdf_pages}",
                    )

                parts: list[str] = []
                segments: list[DocumentSegment] = []
                word_boxes: list[dict[str, Any]] = []
                warnings: list[str] = []
                ocr_pages: list[int] = []
                paragraph_index = 0
                sentence_index = 0
                raw_length = 0

                for page_number, page in enumerate(pdf.pages, start=1):
                    if parts:
                        parts.append("\n")
                        raw_length += 1
                    page_start = raw_length
                    page_text = page.extract_text() or ""
                    if raw_length + len(page_text) > request.config.max_extracted_chars:
                        raise LoaderFailure(
                            DocumentFailureCode.LIMIT_EXCEEDED,
                            (
                                "PDF extracted text exceeds character limit "
                                f"{request.config.max_extracted_chars} on page {page_number}"
                            ),
                        )
                    parts.append(page_text)
                    raw_length += len(page_text)
                    page_end = raw_length
                    segments.append(
                        DocumentSegment(
                            kind="page",
                            page=page_number,
                            raw_start=page_start,
                            raw_end=page_end,
                        )
                    )
                    page_segments, paragraph_index, sentence_index = _text_segments(
                        page_text,
                        base_offset=page_start,
                        page=page_number,
                        paragraph_start=paragraph_index,
                        sentence_start=sentence_index,
                    )
                    segments.extend(page_segments)

                    if not page_text.strip() and bool(getattr(page, "images", ())):
                        ocr_pages.append(page_number)

                    try:
                        words = page.extract_words() or ()
                    except Exception as exc:  # word boxes are supplemental, not fatal
                        warnings.append(
                            f"Page {page_number}: word boxes unavailable ({type(exc).__name__})"
                        )
                        words = ()
                    search_from = 0
                    for word in words:
                        text = str(word.get("text", ""))
                        if not text:
                            continue
                        local_start = page_text.find(text, search_from)
                        if local_start < 0:
                            local_start = page_text.find(text)
                        if local_start < 0:
                            continue
                        local_end = local_start + len(text)
                        search_from = local_end
                        try:
                            box = {
                                "raw_start": page_start + local_start,
                                "raw_end": page_start + local_end,
                                "page": page_number,
                                "x0": float(word["x0"]),
                                "y0": float(word["top"]),
                                "x1": float(word["x1"]),
                                "y1": float(word["bottom"]),
                                "text": text,
                            }
                        except (KeyError, TypeError, ValueError):
                            continue
                        if len(word_boxes) >= request.config.max_pdf_word_boxes:
                            raise LoaderFailure(
                                DocumentFailureCode.LIMIT_EXCEEDED,
                                (
                                    "PDF word boxes exceed limit "
                                    f"{request.config.max_pdf_word_boxes} on page {page_number}"
                                ),
                            )
                        word_boxes.append(box)

                if ocr_pages:
                    pages = ", ".join(str(number) for number in ocr_pages)
                    raise LoaderFailure(
                        DocumentFailureCode.OCR_REQUIRED,
                        f"PDF page(s) contain images but no extractable text: {pages}",
                    )

                return LoadedDocument(
                    raw_text="".join(parts),
                    media_type=self.media_type,
                    segments=tuple(_ordered_segments(segments)),
                    metadata={
                        "format": "pdf",
                        "page_count": len(pdf.pages),
                        "word_boxes": tuple(word_boxes),
                    },
                    warnings=tuple(warnings),
                )
        except LoaderFailure:
            raise
        except Exception as exc:
            if _looks_encrypted(exc):
                raise LoaderFailure(
                    DocumentFailureCode.ENCRYPTED,
                    "PDF is encrypted and requires a password",
                ) from exc
            raise LoaderFailure(
                DocumentFailureCode.MALFORMED,
                f"PDF could not be parsed: {type(exc).__name__}",
            ) from exc


class DocxLoader:
    name = "python-docx"
    version = "1.1.0"
    supported_suffixes = (".docx",)
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def load(self, request: LoadRequest) -> LoadedDocument:
        if request.content.startswith(_OLE_COMPOUND_MAGIC):
            raise LoaderFailure(
                DocumentFailureCode.ENCRYPTED,
                "DOCX is an encrypted Office compound document",
            )
        try:
            _preflight_docx(request)
            from docx import Document as OpenDocument
            from docx.text.paragraph import Paragraph

            package = OpenDocument(io.BytesIO(request.content))
            parts: list[str] = []
            segments: list[DocumentSegment] = []
            paragraph_index = 0
            sentence_index = 0
            raw_length = 0
            table_count = 0

            def separator() -> None:
                nonlocal raw_length
                if parts:
                    parts.append("\n")
                    raw_length += 1

            for block in package.iter_inner_content():
                if isinstance(block, Paragraph):
                    paragraph = block
                    separator()
                    start = raw_length
                    text = paragraph.text
                    _require_extracted_capacity(request, raw_length, text)
                    parts.append(text)
                    raw_length += len(text)
                    block_segments, paragraph_index, sentence_index = _text_segments(
                        text,
                        base_offset=start,
                        paragraph_start=paragraph_index,
                        sentence_start=sentence_index,
                    )
                    segments.extend(block_segments)
                else:
                    table = block
                    table_count += 1
                    for row_number, row in enumerate(table.rows):
                        separator()
                        for column_number, cell in enumerate(row.cells):
                            if column_number:
                                parts.append("\t")
                                raw_length += 1
                            cell_start = raw_length
                            cell_text = "\n".join(paragraph.text for paragraph in cell.paragraphs)
                            _require_extracted_capacity(request, raw_length, cell_text)
                            parts.append(cell_text)
                            raw_length += len(cell_text)
                            segments.append(
                                DocumentSegment(
                                    kind="cell",
                                    label=f"table:{table_count};row:{row_number};column:{column_number}",
                                    raw_start=cell_start,
                                    raw_end=raw_length,
                                )
                            )
                            cell_segments, paragraph_index, sentence_index = _text_segments(
                                cell_text,
                                base_offset=cell_start,
                                paragraph_start=paragraph_index,
                                sentence_start=sentence_index,
                            )
                            segments.extend(cell_segments)

            core = package.core_properties
            metadata = {
                "format": "docx",
                "core_properties": {
                    "author": core.author or None,
                    "created": core.created.isoformat() if core.created else None,
                    "last_modified_by": core.last_modified_by or None,
                    "modified": core.modified.isoformat() if core.modified else None,
                    "subject": core.subject or None,
                    "title": core.title or None,
                },
                "table_count": table_count,
            }
            return LoadedDocument(
                raw_text="".join(parts),
                media_type=self.media_type,
                segments=tuple(_ordered_segments(segments)),
                metadata=metadata,
            )
        except LoaderFailure:
            raise
        except Exception as exc:
            if _looks_encrypted(exc):
                raise LoaderFailure(
                    DocumentFailureCode.ENCRYPTED,
                    "DOCX is encrypted and requires a password",
                ) from exc
            raise LoaderFailure(
                DocumentFailureCode.MALFORMED,
                f"DOCX could not be parsed: {type(exc).__name__}",
            ) from exc


class EmlLoader:
    name = "stdlib-email"
    version = "1.0.0"
    supported_suffixes = (".eml",)
    media_type = "message/rfc822"

    def load(self, request: LoadRequest) -> LoadedDocument:
        try:
            message = BytesParser(policy=policy.default).parsebytes(request.content)
        except Exception as exc:
            raise LoaderFailure(
                DocumentFailureCode.MALFORMED,
                f"EML could not be parsed: {type(exc).__name__}",
            ) from exc

        parts: list[str] = []
        segments: list[DocumentSegment] = []
        raw_length = 0
        headers: dict[str, tuple[str, ...]] = {}
        warnings = [
            f"Email parser defect: {defect.__class__.__name__}" for defect in message.defects
        ]

        def append_separator() -> None:
            nonlocal raw_length
            if parts:
                parts.append("\n")
                raw_length += 1

        for header_name in _SEMANTIC_HEADERS:
            values = tuple(str(value) for value in message.get_all(header_name, ()))
            headers[header_name] = values
            for value in values:
                append_separator()
                label = f"{header_name.title()}: "
                parts.append(label)
                raw_length += len(label)
                value_start = raw_length
                parts.append(value)
                raw_length += len(value)
                segments.append(
                    DocumentSegment(
                        kind="field",
                        field=header_name,
                        raw_start=value_start,
                        raw_end=raw_length,
                    )
                )

        plain_bodies: list[str] = []
        html_bodies: list[str] = []
        attachments: list[dict[str, Any]] = []
        attachment_content: list[tuple[str, str | None, str]] = []

        leaf_parts = tuple(message.walk()) if message.is_multipart() else (message,)
        for part in leaf_parts:
            if part.is_multipart():
                continue
            disposition = part.get_content_disposition()
            filename = part.get_filename()
            content_type = part.get_content_type().lower()
            is_attachment = disposition == "attachment" or filename is not None
            payload = _message_part_bytes(part)
            if is_attachment:
                included_text = content_type.startswith("text/")
                text = ""
                if included_text:
                    text, decode_warning = _decode_message_text(part, payload)
                    if decode_warning:
                        warnings.append(decode_warning)
                    if content_type == "text/html":
                        text = _html_to_text(text)
                attachment_name = str(filename) if filename is not None else None
                attachments.append(
                    {
                        "filename": attachment_name,
                        "content_type": content_type,
                        "size_bytes": len(payload),
                        "sha256": hashlib.sha256(payload).hexdigest(),
                        "included_text": included_text,
                    }
                )
                attachment_content.append((content_type, attachment_name, text))
                continue

            if content_type not in {"text/plain", "text/html"}:
                continue
            text, decode_warning = _decode_message_text(part, payload)
            if decode_warning:
                warnings.append(decode_warning)
            if content_type == "text/plain":
                plain_bodies.append(text)
            else:
                html_bodies.append(_html_to_text(text))

        chosen_bodies = plain_bodies if plain_bodies else html_bodies
        if chosen_bodies:
            append_separator()
            body_start = raw_length
            body = "\n".join(chosen_bodies)
            parts.append(body)
            raw_length += len(body)
            segments.append(
                DocumentSegment(
                    kind="field",
                    field="body",
                    raw_start=body_start,
                    raw_end=raw_length,
                )
            )
            body_segments, _, _ = _text_segments(body, base_offset=body_start)
            segments.extend(body_segments)

        for content_type, filename, text in attachment_content:
            append_separator()
            attachment_start = raw_length
            if text:
                parts.append(text)
                raw_length += len(text)
            segments.append(
                DocumentSegment(
                    kind="attachment",
                    label=filename or content_type,
                    raw_start=attachment_start,
                    raw_end=raw_length,
                )
            )

        return LoadedDocument(
            raw_text="".join(parts),
            media_type=self.media_type,
            segments=tuple(_ordered_segments(segments)),
            metadata={
                "format": "eml",
                "headers": headers,
                "attachments": tuple(attachments),
            },
            warnings=tuple(warnings),
        )


def _preflight_docx(request: LoadRequest) -> None:
    """Reject dangerous OPC ZIP structures before ``python-docx`` expands them."""

    try:
        with zipfile.ZipFile(io.BytesIO(request.content)) as archive:
            members = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise LoaderFailure(
            DocumentFailureCode.MALFORMED, "DOCX is not a valid ZIP package"
        ) from exc

    if len(members) > request.config.max_docx_members:
        raise LoaderFailure(
            DocumentFailureCode.LIMIT_EXCEEDED,
            (
                f"DOCX has {len(members)} archive members; limit is "
                f"{request.config.max_docx_members}"
            ),
        )

    total_uncompressed = 0
    names: set[str] = set()
    for member in members:
        if member.filename in names:
            raise LoaderFailure(
                DocumentFailureCode.MALFORMED,
                f"DOCX contains a duplicate archive member: {member.filename!r}",
            )
        names.add(member.filename)
        if _unsafe_archive_name(member.filename):
            raise LoaderFailure(
                DocumentFailureCode.MALFORMED,
                f"DOCX contains an unsafe archive member name: {member.filename!r}",
            )
        if member.flag_bits & 0x1:
            raise LoaderFailure(
                DocumentFailureCode.ENCRYPTED,
                f"DOCX archive member is encrypted: {member.filename!r}",
            )
        if member.file_size > request.config.max_docx_member_bytes:
            raise LoaderFailure(
                DocumentFailureCode.LIMIT_EXCEEDED,
                (
                    f"DOCX member {member.filename!r} is {member.file_size} bytes; limit is "
                    f"{request.config.max_docx_member_bytes}"
                ),
            )
        total_uncompressed += member.file_size
        if total_uncompressed > request.config.max_docx_uncompressed_bytes:
            raise LoaderFailure(
                DocumentFailureCode.LIMIT_EXCEEDED,
                (
                    "DOCX uncompressed archive bytes exceed limit "
                    f"{request.config.max_docx_uncompressed_bytes}"
                ),
            )
        if member.file_size:
            ratio = member.file_size / max(member.compress_size, 1)
            if ratio > request.config.max_docx_compression_ratio:
                raise LoaderFailure(
                    DocumentFailureCode.LIMIT_EXCEEDED,
                    (
                        f"DOCX member {member.filename!r} has compression ratio {ratio:.1f}; "
                        f"limit is {request.config.max_docx_compression_ratio:g}"
                    ),
                )


def _unsafe_archive_name(name: str) -> bool:
    if not name or "\\" in name or name.startswith(("/", "\\")):
        return True
    path = PurePosixPath(name)
    return path.is_absolute() or ".." in path.parts or bool(path.parts and ":" in path.parts[0])


def _require_extracted_capacity(request: LoadRequest, current_length: int, text: str) -> None:
    if current_length + len(text) > request.config.max_extracted_chars:
        raise LoaderFailure(
            DocumentFailureCode.LIMIT_EXCEEDED,
            f"DOCX extracted text exceeds character limit {request.config.max_extracted_chars}",
        )


def _decode_source_text(content: bytes, filename: str) -> str:
    if content.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    elif content.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    else:
        encoding = "utf-8"
    try:
        return content.decode(encoding, errors="strict")
    except (UnicodeDecodeError, LookupError) as exc:
        raise LoaderFailure(
            DocumentFailureCode.MALFORMED,
            f"{filename} is not valid {encoding} text",
        ) from exc


def _text_segments(
    text: str,
    *,
    base_offset: int = 0,
    page: int | None = None,
    paragraph_start: int = 0,
    sentence_start: int = 0,
) -> tuple[list[DocumentSegment], int, int]:
    segments: list[DocumentSegment] = []
    paragraphs = _paragraph_ranges(text)
    paragraph_index = paragraph_start
    sentence_index = sentence_start
    for local_start, local_end in paragraphs:
        segments.append(
            DocumentSegment(
                kind="paragraph",
                page=page,
                paragraph=paragraph_index,
                raw_start=base_offset + local_start,
                raw_end=base_offset + local_end,
            )
        )
        paragraph_text = text[local_start:local_end]
        for match in re.finditer(r"[^.!?]+(?:[.!?]+|$)", paragraph_text):
            candidate = match.group(0)
            if not candidate.strip():
                continue
            leading = len(candidate) - len(candidate.lstrip())
            trailing = len(candidate.rstrip())
            sentence_local_start = local_start + match.start() + leading
            sentence_local_end = local_start + match.start() + trailing
            segments.append(
                DocumentSegment(
                    kind="sentence",
                    page=page,
                    paragraph=paragraph_index,
                    sentence=sentence_index,
                    raw_start=base_offset + sentence_local_start,
                    raw_end=base_offset + sentence_local_end,
                )
            )
            sentence_index += 1
        paragraph_index += 1
    return segments, paragraph_index, sentence_index


def _paragraph_ranges(text: str) -> list[tuple[int, int]]:
    ranges: list[tuple[int, int]] = []
    offset = 0
    paragraph_start: int | None = None
    paragraph_end = 0
    for line in text.splitlines(keepends=True):
        content = line.rstrip("\r\n")
        content_start = offset
        content_end = offset + len(content)
        if content.strip():
            if paragraph_start is None:
                leading = len(content) - len(content.lstrip())
                paragraph_start = content_start + leading
            paragraph_end = content_end - (len(content) - len(content.rstrip()))
        elif paragraph_start is not None:
            ranges.append((paragraph_start, paragraph_end))
            paragraph_start = None
        offset += len(line)
    if offset < len(text):  # splitlines omits nothing, kept for defensive alternate implementations
        offset = len(text)
    if paragraph_start is not None:
        ranges.append((paragraph_start, paragraph_end))
    if not text:
        return ranges
    if not text.splitlines(keepends=True) and text.strip():
        start = len(text) - len(text.lstrip())
        end = len(text.rstrip())
        ranges.append((start, end))
    return ranges


def _message_part_bytes(part: Message) -> bytes:
    payload = part.get_payload(decode=True)
    if isinstance(payload, bytes):
        return payload
    raw_payload = part.get_payload()
    if isinstance(raw_payload, str):
        return raw_payload.encode("utf-8", errors="surrogatepass")
    return part.as_bytes(policy=policy.default)


def _decode_message_text(part: Message, payload: bytes) -> tuple[str, str | None]:
    charset = part.get_content_charset() or "utf-8"
    try:
        return payload.decode(charset, errors="strict"), None
    except (LookupError, UnicodeDecodeError):
        return (
            payload.decode("utf-8", errors="replace"),
            f"Email {part.get_content_type()} part used replacement decoding",
        )


def _html_to_text(html: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for element in soup(("script", "style", "template")):
        element.decompose()
    lines = (line.strip() for line in soup.get_text("\n").splitlines())
    return "\n".join(line for line in lines if line)


def _ordered_segments(segments: list[DocumentSegment]) -> list[DocumentSegment]:
    kind_order = {
        "page": 0,
        "paragraph": 1,
        "sentence": 2,
        "field": 3,
        "attachment": 4,
        "sheet": 5,
        "cell": 6,
    }
    return sorted(
        segments,
        key=lambda segment: (
            segment.raw_start,
            -segment.raw_end,
            kind_order[segment.kind],
            segment.field or "",
            segment.label or "",
        ),
    )


def _looks_encrypted(exc: Exception) -> bool:
    name = type(exc).__name__.casefold()
    message = str(exc).casefold()
    return "password" in name or "password" in message or "encrypt" in message


__all__ = ["DocxLoader", "EmlLoader", "MarkdownLoader", "PdfLoader", "TextLoader"]
